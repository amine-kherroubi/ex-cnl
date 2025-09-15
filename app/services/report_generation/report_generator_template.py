from __future__ import annotations

# Standard library imports
from abc import ABC, abstractmethod
from pathlib import Path
from typing import TYPE_CHECKING, Callable
from logging import Logger

# Third-party imports
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

# Local application imports
from app.data.data_repository import DataRepository
from app.services.report_generation.business_values.programmes import (
    get_programmes_dataframe,
)
from app.services.file_io.file_io_service import FileIOService
from app.services.report_generation.models.report_context import (
    ReportContext,
)
from app.utils.exceptions import QueryExecutionError
from app.utils.logging_setup import get_logger

if TYPE_CHECKING:
    from app.services.report_generation.report_specification_registry import (
        ReportSpecification,
    )


class ReportGenerator(ABC):
    __slots__ = (
        "_file_io_service",
        "_data_repository",
        "_report_specification",
        "_report_context",
        "_workbook",
        "_logger",
    )

    def __init__(
        self,
        file_io_service: FileIOService,
        data_repository: DataRepository,
        report_specification: ReportSpecification,
        report_context: ReportContext,
    ) -> None:
        self._logger: Logger = get_logger(
            f"app.services.report_generation.report_generator_template"
        )
        self._logger.debug(f"Initializing {self.__class__.__name__}")

        self._file_io_service: FileIOService = file_io_service
        self._data_repository: DataRepository = data_repository
        self._report_specification: ReportSpecification = report_specification
        self._report_context: ReportContext = report_context
        self._workbook: Workbook | None = None

        self._logger.info(
            f"Generator initialized for report: {report_specification.display_name}"
        )
        self._logger.debug(
            f"Report context: wilaya={report_context.wilaya.value}, date={report_context.report_date}"
        )

    def generate(
        self,
        source_file_paths: dict[str, Path],
        output_directory_path: Path,
    ) -> Path:
        self._logger.info("Starting report generation process")

        try:
            # Step 1: Loading data into in-memory database
            self._logger.debug("Step 1: Loading data into database")
            self._load_data_into_db(source_file_paths)
            self._logger.info("Data successfully loaded into database")

            # Step 2: Creating reference tables
            self._logger.debug("Step 2: Creating reference tables")
            self._create_predefined_tables()
            self._logger.info("Program reference table created successfully")

            # Step 3: Executing queries
            self._logger.debug("Step 3: Executing queries")
            query_results: dict[str, pd.DataFrame] = self._execute_queries()
            self._logger.info(f"Successfully executed {len(query_results)} queries")

            # Step 4: Creating workbook and main worksheet
            self._logger.debug("Step 4: Creating Excel workbook")
            self._workbook = Workbook()
            self._workbook.remove(self._workbook.active)  # type: ignore
            sheet: Worksheet = self._workbook.create_sheet(
                self._report_specification.display_name
            )
            self._logger.info(
                f"Workbook created with sheet: {self._report_specification.display_name}"
            )

            # Step 5: Adding header (subclass responsibility)
            self._logger.debug("Step 5: Adding report header")
            self._add_header(sheet)
            self._logger.debug("Header added successfully")

            # Step 6: Building tables (subclass responsibility)
            self._logger.debug("Step 6: Building main table")
            self._add_tables(sheet, query_results)
            self._logger.debug("Main table built successfully")

            # Step 7: Adding footer (subclass responsibility)
            self._logger.debug("Step 7: Adding report footer")
            self._add_footer(sheet)
            self._logger.debug("Footer added successfully")

            # Step 8: Final formatting (subclass responsibility)
            self._logger.debug("Step 8: Applying final formatting")
            self._finalize_formatting(sheet)
            self._logger.debug("Final formatting applied successfully")

            # Step 9: Generate actual (formatted) filename
            output_file_path: Path = (
                output_directory_path / self._generate_output_filename()
            )

            # Step 10: Saving report
            self._logger.debug("Step 10: Saving report")
            self._save_report(output_file_path)
            self._logger.info("Report generation completed successfully")

            return output_file_path

        except Exception as error:
            self._logger.exception(f"Report generation failed: {error}")
            raise

    def _load_data_into_db(self, source_file_paths: dict[str, Path]) -> None:
        self._logger.debug("Loading files into database views")

        for table_name, file_path in source_file_paths.items():
            self._logger.debug(f"Loading file '{file_path}' into view '{table_name}'")
            try:
                df: pd.DataFrame = self._file_io_service.load_data_from_file(file_path)

                original_columns: list[str] = list(df.columns)
                df.columns = [column.strip() for column in df.columns]
                cleaned_columns: list[str] = list(df.columns)

                if original_columns != cleaned_columns:
                    self._logger.debug(f"Column names cleaned for view '{table_name}'")

                self._data_repository.create_table_from_dataframe(table_name, df)
                self._logger.info(
                    f"View '{table_name}' loaded successfully: {len(df)} rows"
                )

                print(self._data_repository.summarize(table_name))

            except Exception as error:
                self._logger.error(
                    f"Failed to load file '{file_path}' into view '{table_name}': {error}"
                )
                raise

    def _create_predefined_tables(self) -> None:
        self._logger.debug("Creating reference tables")

        predefined_tables: dict[str, Callable[[], pd.DataFrame]] = {
            "programmes": get_programmes_dataframe,
        }

        for table_name, dataframe_factory in predefined_tables.items():
            try:
                self._logger.debug(f"Creating reference table '{table_name}'")

                df: pd.DataFrame = dataframe_factory()
                self._data_repository.create_table_from_dataframe(table_name, df)

                rows, cols = df.shape
                self._logger.info(
                    f"Reference table '{table_name}' created: {rows} rows and {cols} columns"
                )
                self._logger.debug(f"Columns for '{table_name}': {list(df.columns)}")

            except Exception as error:
                self._logger.exception(
                    f"Failed to create reference table '{table_name}': {error}"
                )
                raise

    def _execute_queries(self) -> dict[str, pd.DataFrame]:
        self._logger.debug("Executing report queries")

        results: dict[str, pd.DataFrame] = {}
        query_count: int = len(self._report_specification.queries)
        self._logger.debug(f"Preparing to execute {query_count} queries")

        for query_name, query_template in self._report_specification.queries.items():
            self._logger.debug(f"Executing query '{query_name}'")
            try:
                formatted_query: str = self._format_query_with_context(query_template)
                self._logger.debug(
                    f"Query '{query_name}' formatted with report context"
                )

                result_df: pd.DataFrame = self._data_repository.execute(formatted_query)
                results[query_name] = result_df

                self._logger.info(
                    f"Query '{query_name}' returned {len(result_df)} rows"
                )

            except Exception as error:
                self._logger.exception(f"Query '{query_name}' failed: {error}")
                raise QueryExecutionError(
                    f"Required query '{query_name}' failed", error
                )

        self._logger.info(f"All {query_count} queries executed successfully")
        return results

    def _format_query_with_context(self, query_template: str) -> str:
        self._logger.debug("Formatting query template with report context")

        formatted_query: str = query_template

        if self._report_context.month:
            month_number: int = self._report_context.month.number
            year: int = self._report_context.year

            formatted_query = formatted_query.replace(
                "{month_number:02d}", f"{month_number:02d}"
            )
            formatted_query = formatted_query.replace(
                "{month_number}", str(month_number)
            )
            formatted_query = formatted_query.replace("{year}", str(year))

            self._logger.debug(
                f"Placeholders replaced with: month_number={month_number:02d}, year={year}"
            )

        formatted_query = formatted_query.replace(
            "{year}", str(self._report_context.year)
        )

        self._logger.debug("Query formatting completed")
        return formatted_query

    @abstractmethod
    def _add_header(self, sheet: Worksheet) -> None: ...

    @abstractmethod
    def _add_tables(
        self, sheet: Worksheet, query_results: dict[str, pd.DataFrame]
    ) -> None: ...

    @abstractmethod
    def _add_footer(self, sheet: Worksheet) -> None: ...

    @abstractmethod
    def _finalize_formatting(self, sheet: Worksheet) -> None: ...

    def _generate_output_filename(self) -> str:
        from app.utils.date_formatting import DateFormatter

        output_filename: str = self._report_specification.output_filename
        self._logger.debug(f"Output filename template: {output_filename}")

        output_filename = output_filename.replace(
            "{wilaya}", self._report_context.wilaya.value
        )
        output_filename = output_filename.replace(
            "{date}",
            DateFormatter.to_french_filename_date(self._report_context.report_date),
        )

        if not output_filename.endswith(".xlsx"):
            output_filename += ".xlsx"

        self._logger.debug(f"Generated filename: {output_filename}")
        return output_filename

    def _save_report(self, output_file_path: Path) -> None:
        self._logger.debug(f"Saving report to: {output_file_path}")

        if not self._workbook:
            error_msg: str = "No workbook created"
            self._logger.error(error_msg)
            raise ValueError(error_msg)

        try:
            self._file_io_service.save_data_to_file(
                data=self._workbook,
                output_file_path=output_file_path,
            )
            self._logger.info(f"Report saved successfully: {output_file_path}")
        except Exception as error:
            self._logger.exception(f"Failed to save report: {error}")
            raise
