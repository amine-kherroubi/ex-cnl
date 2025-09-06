from __future__ import annotations

# Standard library imports
from abc import ABC, abstractmethod
from typing import Any

# Third-party imports
import pandas
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

# Local application imports
from app.data.data_repository import DataRepository
from app.services.document_generation.documents_registry import DocumentSpecification
from app.services.file_storage.file_storage_service import FileStorageService


class DocumentGenerator(ABC):  # Template Method pattern
    __slots__ = (
        "_storage_service",
        "_data_repository",
        "_document_definition",
        "_workbook",
    )

    def __init__(
        self, storage_service: FileStorageService, data_repository: DataRepository
    ) -> None:
        self._storage_service: FileStorageService = storage_service
        self._data_repository: DataRepository = data_repository
        self._document_definition: DocumentSpecification
        self._workbook: Workbook | None = None

    def generate(self) -> None:
        # Step 1: Validate required files
        files_matching_patterns: dict[str, str] = self._validate_required_files()

        # Step 2: Load data into in-memory database
        self._load_data_into_db(files_matching_patterns)

        # Step 3: Execute queries
        query_results: dict[str, pandas.DataFrame] = self._execute_queries()

        # Step 4: Create workbook and main sheet
        self._workbook = Workbook()
        self._workbook.remove(self._workbook.active)  # type: ignore
        sheet: Any = self._workbook.create_sheet(self._document_definition.display_name)

        # Step 5: Add header (subclass responsibility)
        self._add_header(sheet)

        # Step 6: Build main table (subclass responsibility)
        self._add_table(sheet, query_results)

        # Step 7: Add footer (subclass responsibility)
        self._add_footer(sheet)

        # Step 8: Final formatting (subclass responsibility)
        self._finalize_formatting(sheet)

        # Step 9: Save document
        self._save_document()

    def _validate_required_files(self) -> dict[str, str]:
        if not self._document_definition:
            raise ValueError("Document specification not set")

        missing_patterns: list[str] = []
        files_mapping: dict[str, str] = {}

        for pattern, view_name in self._document_definition.required_files.items():
            filename: str | None = self._storage_service.find_filename_matching_pattern(
                pattern
            )
            if filename is None:
                missing_patterns.append(pattern)
            else:
                files_mapping[view_name] = filename

        if missing_patterns:
            raise FileNotFoundError(
                f"Missing required files matching patterns: {missing_patterns}"
            )

        return files_mapping

    def _load_data_into_db(self, files_mapping: dict[str, str]) -> None:
        for view_name, filename in files_mapping.items():
            df: pandas.DataFrame = self._storage_service.load_data_from_file(filename)
            self._data_repository.create_view_from_dataframe(view_name, df)

    def _execute_queries(self) -> dict[str, pandas.DataFrame]:
        if not self._document_definition:
            raise ValueError("Document specification not set")

        results: dict[str, pandas.DataFrame] = {}
        for query_name, query in self._document_definition.queries.items():
            try:
                results[query_name] = self._data_repository.execute(query)
            except Exception as e:
                print(f"Query '{query_name}' failed: {e}")
                results[query_name] = pandas.DataFrame()
        return results

    @abstractmethod
    def _add_header(self, sheet: Worksheet) -> None: ...

    @abstractmethod
    def _add_table(
        self, sheet: Worksheet, query_results: dict[str, pandas.DataFrame]
    ) -> None: ...

    @abstractmethod
    def _add_footer(self, sheet: Worksheet) -> None: ...

    @abstractmethod
    def _finalize_formatting(self, sheet: Worksheet) -> None: ...

    def _save_document(self) -> None:
        if self._workbook:
            self._storage_service.save_data_to_file(
                data=self._workbook,
                output_filename=self._document_definition.output_filename,
            )
